#!/usr/bin/env python3
"""
Convert KAALADRISHTI_AND_KAALFORM_SPECIFICATION.md into a professional .docx Word Document.
Features:
- Professional styling (Executive palette: Deep Navy, Steel Blue, Charcoal)
- Title & subtitle headers
- Custom styled tables with bold headers & alternating rows
- Code blocks in Courier / Consolas with background shading
- Bullet lists & blockquotes
"""

import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)

def create_document():
    doc = Document()

    # Page Margins (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Style definitions
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    normal_style.paragraph_format.line_spacing = 1.18
    normal_style.paragraph_format.space_after = Pt(6)

    md_path = os.path.join(os.path.dirname(__file__), "..", "docs", "KAALADRISHTI_AND_KAALFORM_SPECIFICATION.md")
    output_docx = os.path.join(os.path.dirname(__file__), "..", "docs", "KAALADRISHTI_AND_KAALFORM_SPECIFICATION.docx")

    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_code_block = False
    code_block_lines = []
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        
        # Determine max cols
        num_cols = max(len(r) for r in table_rows)
        table = doc.add_table(rows=len(table_rows), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        for row_idx, row_data in enumerate(table_rows):
            is_header = (row_idx == 0)
            row = table.rows[row_idx]
            for col_idx in range(num_cols):
                cell = row.cells[col_idx]
                val = row_data[col_idx] if col_idx < len(row_data) else ""
                cell.text = val
                
                # Format cell text
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                if p.runs:
                    run = p.runs[0]
                    run.font.name = 'Calibri'
                    if is_header:
                        run.font.bold = True
                        run.font.size = Pt(10)
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    else:
                        run.font.size = Pt(9.5)
                        run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

                # Background shading
                if is_header:
                    set_cell_background(cell, "1E3A8A") # Navy
                elif row_idx % 2 == 1:
                    set_cell_background(cell, "F8FAFC") # Soft gray
                else:
                    set_cell_background(cell, "FFFFFF")

                set_cell_margins(cell, top=120, bottom=120, left=150, right=150)

        # Space after table
        empty_p = doc.add_paragraph()
        empty_p.paragraph_format.space_after = Pt(6)
        table_rows = []
        in_table = False

    def flush_code_block():
        nonlocal in_code_block, code_block_lines
        if not code_block_lines:
            in_code_block = False
            return
        code_text = "".join(code_block_lines)
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.right_indent = Inches(0.2)
        run = p.add_run(code_text.rstrip())
        run.font.name = 'Consolas'
        run.font.size = Pt(9.0)
        run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        
        # Border / background for code paragraph
        pPr = p._p.get_or_add_pPr()
        pBdr = parse_xml(
            f'<w:pBdr {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="4" w:color="CBD5E1"/>'
            f'<w:left w:val="single" w:sz="18" w:space="8" w:color="3B82F6"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="4" w:color="CBD5E1"/>'
            f'<w:right w:val="single" w:sz="6" w:space="4" w:color="CBD5E1"/>'
            f'</w:pBdr>'
        )
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
        pPr.append(pBdr)
        pPr.append(shd)

        code_block_lines = []
        in_code_block = False

    for line in lines:
        stripped = line.strip()

        # Handle code blocks
        if stripped.startswith("```"):
            if in_code_block:
                flush_code_block()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
            continue

        if in_code_block:
            code_block_lines.append(line)
            continue

        # Handle Markdown Tables
        if "|" in line and (line.startswith("|") or stripped.startswith("|")):
            if re.match(r'^\s*\|?\s*[-:]+[-| :]*\|?\s*$', stripped):
                # Table delimiter line, ignore
                continue
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            table_rows.append(cells)
            in_table = True
            continue
        elif in_table:
            flush_table()

        # Empty line
        if not stripped:
            continue

        # Horizontal Rule
        if stripped in ["---", "***", "___"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            pPr = p._p.get_or_add_pPr()
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="1" w:color="CBD5E1"/></w:pBdr>')
            pPr.append(pBdr)
            continue

        # Headers
        if stripped.startswith("# "):
            h_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(h_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A) # Dark Navy
            continue

        if stripped.startswith("## "):
            h_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(h_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Deep Blue
            continue

        if stripped.startswith("### "):
            h_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(h_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB) # Slate Blue
            continue

        if stripped.startswith("#### "):
            h_text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(h_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(11.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            continue

        # Bullet List
        if stripped.startswith("* ") or stripped.startswith("- "):
            bullet_text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            add_formatted_text(p, bullet_text)
            continue

        # Numbered List
        m_num = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if m_num:
            num_text = m_num.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            add_formatted_text(p, num_text)
            continue

        # Blockquote
        if stripped.startswith("> "):
            quote_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            pPr = p._p.get_or_add_pPr()
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="8" w:color="0284C7"/></w:pBdr>')
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F0F9FF"/>')
            pPr.append(pBdr)
            pPr.append(shd)
            add_formatted_text(p, quote_text, italic_all=True)
            continue

        # Regular Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        add_formatted_text(p, stripped)

    if in_table:
        flush_table()
    if in_code_block:
        flush_code_block()

    doc.save(output_docx)
    print(f"Successfully generated DOCX at: {output_docx}")

def add_formatted_text(paragraph, text, italic_all=False):
    """Parses inline bold (**bold**), code (`code`), and italics."""
    pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)')
    parts = pattern.split(text)

    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            run.font.name = 'Calibri'
            if italic_all:
                run.italic = True
        elif part.startswith("`") and part.endswith("`") and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0xBE, 0x18, 0x5D) # Subtle magenta code
            if italic_all:
                run.italic = True
        elif part.startswith("*") and part.endswith("*") and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
            run.font.name = 'Calibri'
        else:
            run = paragraph.add_run(part)
            run.font.name = 'Calibri'
            if italic_all:
                run.italic = True

if __name__ == "__main__":
    create_document()
